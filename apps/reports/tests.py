"""
Tests for the reports app: export service, analytics and views.
"""
from decimal import Decimal
import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
import openpyxl

from apps.expenses.models import Expense, ExpenseCategory, CategoryKind
from apps.fuel.models import FuelType, Refueling
from apps.maintenance.models import Maintenance, MaintenanceType, ServiceType
from apps.vehicles.models import Vehicle

from .services.analytics_service import VehicleComparisonService
from .services.export_service import ExportService

User = get_user_model()


class ExportServiceTests(TestCase):
    """Unit tests for ExportService (CSV and Excel)."""

    def setUp(self):
        self.user = User.objects.create_user(
            email="export_user@example.com",
            password="testpassword123",
        )
        self.vehicle = Vehicle.objects.create(
            owner_primary=self.user,
            name="C3 Teste",
            plate="ABC1234",
            year=2020,
            initial_odometer=10000,
            current_odometer_cache=10500,
        )
        self.fuel_type = FuelType.objects.create(
            user=self.user,
            name="Gasolina Comum",
            code="GAS",
        )
        self.refueling = Refueling.objects.create(
            vehicle=self.vehicle,
            occurred_at=timezone.now().date(),
            odometer=10500,
            fuel_type=self.fuel_type,
            liters=Decimal("40.000"),
            price_per_liter=Decimal("5.500"),
            total_amount=Decimal("220.00"),
            is_full_tank=True,
            station_name="Posto Shell",
        )
        self.category = ExpenseCategory.objects.create(
            user=self.user,
            name="Estacionamento",
            slug="estacionamento",
            kind=CategoryKind.VARIABLE,
        )
        self.expense = Expense.objects.create(
            vehicle=self.vehicle,
            category=self.category,
            occurred_at=timezone.now().date(),
            amount=Decimal("30.00"),
            vendor_name="Shopping",
        )

    def test_export_fuel_csv(self):
        """CSV export returns status 200, utf-8-sig encoding and semicolon delimiter."""
        response = ExportService.export(
            user=self.user,
            entity="fuel",
            export_format="csv",
            vehicle=self.vehicle,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("attachment;", response["Content-Disposition"])

        content = response.content.decode("utf-8-sig")
        lines = content.strip().split("\r\n")
        self.assertGreaterEqual(len(lines), 2)
        # Semicolon in header
        self.assertIn("Data;Veículo;Odômetro (km)", lines[0])
        # Data present
        self.assertIn("C3 Teste", lines[1])
        self.assertIn("Posto Shell", lines[1])

    def test_export_consolidated_excel(self):
        """Excel export creates a multi-sheet workbook with openpyxl."""
        response = ExportService.export(
            user=self.user,
            entity="consolidated",
            export_format="xlsx",
            vehicle=self.vehicle,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response["Content-Type"])

        # Inspect workbook
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet_names = wb.sheetnames
        self.assertIn("Abastecimentos", sheet_names)
        self.assertIn("Despesas", sheet_names)
        self.assertIn("Manutenções", sheet_names)
        self.assertIn("Viagens", sheet_names)
        self.assertIn("Lembretes", sheet_names)

        fuel_ws = wb["Abastecimentos"]
        self.assertEqual(fuel_ws.cell(row=1, column=1).value, "Data")
        self.assertEqual(fuel_ws.cell(row=2, column=2).value, "C3 Teste")


class AnalyticsServiceTests(TestCase):
    """Unit tests for VehicleComparisonService."""

    def setUp(self):
        self.user = User.objects.create_user(
            email="analytics_user@example.com",
            password="testpassword123",
        )
        self.v1 = Vehicle.objects.create(
            owner_primary=self.user,
            name="Carro 1",
            year=2020,
            initial_odometer=10000,
            current_odometer_cache=10500,
        )
        self.v2 = Vehicle.objects.create(
            owner_primary=self.user,
            name="Carro 2",
            year=2021,
            initial_odometer=20000,
            current_odometer_cache=21000,
        )
        # Refuelings
        Refueling.objects.create(
            vehicle=self.v1,
            occurred_at=timezone.now().date(),
            odometer=10000,
            liters=Decimal("50.00"),
            total_amount=Decimal("250.00"),
            is_full_tank=True,
        )
        r1_2 = Refueling.objects.create(
            vehicle=self.v1,
            occurred_at=timezone.now().date(),
            odometer=10500,
            liters=Decimal("40.00"),
            total_amount=Decimal("200.00"),
            is_full_tank=True,
            consumption_km_l=Decimal("12.50"),
        )
        # v1 distance: 500km, consumption: 500 / 40 = 12.5 km/l, total fuel: 450.00

    def test_comparison_metrics(self):
        """Comparison service computes metrics and highlights best performers."""
        data = VehicleComparisonService.get_comparison_data(self.user)
        v_data = data["vehicles_data"]
        self.assertEqual(len(v_data), 2)

        v1_metrics = next(item for item in v_data if item["vehicle"] == self.v1)
        self.assertEqual(v1_metrics["distance_km"], 500)
        self.assertEqual(v1_metrics["fuel_cost"], Decimal("450.00"))
        # 450 / 500 = 0.90 R$/km
        self.assertEqual(v1_metrics["cost_per_km"], Decimal("0.90"))
        self.assertEqual(v1_metrics["avg_consumption"], 12.5)


class ReportsViewsTests(TestCase):
    """View tests for reports app."""

    def setUp(self):
        self.user = User.objects.create_user(
            email="reports_views@example.com",
            password="testpassword123",
        )
        self.vehicle = Vehicle.objects.create(
            owner_primary=self.user,
            name="Meu Carro",
            year=2022,
        )
        self.client.login(email=self.user.email, password="testpassword123")

    def test_reports_hub_view(self):
        res = self.client.get(reverse("reports:hub"))
        self.assertEqual(res.status_code, 200)
        self.assertTemplateUsed(res, "reports/reports_hub.html")

    def test_vehicle_comparison_view(self):
        res = self.client.get(reverse("reports:comparison"))
        self.assertEqual(res.status_code, 200)
        self.assertTemplateUsed(res, "reports/vehicle_comparison.html")

    def test_export_data_view_render_and_post(self):
        res_get = self.client.get(reverse("reports:export"))
        self.assertEqual(res_get.status_code, 200)

        res_post = self.client.post(reverse("reports:export"), {
            "entity": "fuel",
            "format": "csv",
        })
        self.assertEqual(res_post.status_code, 200)
        self.assertIn("text/csv", res_post["Content-Type"])

    def test_quick_export_view(self):
        res = self.client.get(reverse("reports:quick_export", kwargs={"entity": "expenses"}))
        self.assertEqual(res.status_code, 200)
        self.assertIn("text/csv", res["Content-Type"])


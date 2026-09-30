"""
Export service supporting CSV (UTF-8 BOM, semicolon delimiter) and Excel (.xlsx).
"""
import csv
import io
from decimal import Decimal
from typing import Optional

from django.http import HttpResponse
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.expenses.models import Expense
from apps.fuel.models import Refueling
from apps.maintenance.models import Maintenance
from apps.reminders.models import Reminder
from apps.trips.models import Trip


class ExportService:
    """Handles data extraction and file formatting for CSV and Excel."""

    HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ROW_FONT = Font(name="Calibri", size=11)

    @classmethod
    def export(
        cls,
        user,
        entity: str,
        export_format: str = "csv",
        vehicle=None,
        date_start=None,
        date_end=None,
    ) -> HttpResponse:
        """
        Main entry point for generating export responses.
        """
        export_format = export_format.lower()
        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        vehicle_slug = f"_{vehicle.name.lower().replace(' ', '_')}" if vehicle else ""
        filename = f"car_management_{entity}{vehicle_slug}_{timestamp}.{export_format}"

        if export_format == "xlsx":
            return cls._export_excel(user, entity, filename, vehicle, date_start, date_end)
        return cls._export_csv(user, entity, filename, vehicle, date_start, date_end)

    # -------------------------------------------------------------------------
    # CSV Export
    # -------------------------------------------------------------------------
    @classmethod
    def _export_csv(cls, user, entity, filename, vehicle, date_start, date_end) -> HttpResponse:
        response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

        writer = csv.writer(response, delimiter=";", quoting=csv.QUOTE_MINIMAL)

        if entity == "fuel":
            cls._write_fuel_csv(writer, user, vehicle, date_start, date_end)
        elif entity == "expenses":
            cls._write_expenses_csv(writer, user, vehicle, date_start, date_end)
        elif entity == "maintenance":
            cls._write_maintenance_csv(writer, user, vehicle, date_start, date_end)
        elif entity == "trips":
            cls._write_trips_csv(writer, user, vehicle, date_start, date_end)
        elif entity == "reminders":
            cls._write_reminders_csv(writer, user, vehicle)
        else:
            # Consolidated
            cls._write_consolidated_csv(writer, user, vehicle, date_start, date_end)

        return response

    @classmethod
    def _write_fuel_csv(cls, writer, user, vehicle, date_start, date_end):
        writer.writerow([
            "Data", "Veículo", "Odômetro (km)", "Combustível", "Preço/L (R$)",
            "Volume (L)", "Valor Total (R$)", "Tanque Cheio", "Consumo (km/l)", "Posto", "Observações"
        ])
        qs = Refueling.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "fuel_type")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(occurred_at__gte=date_start)
        if date_end:
            qs = qs.filter(occurred_at__lte=date_end)

        for r in qs.order_by("-occurred_at", "-odometer"):
            writer.writerow([
                r.occurred_at.strftime("%d/%m/%Y"),
                r.vehicle.name,
                r.odometer,
                r.fuel_type.name if r.fuel_type else "",
                f"{r.price_per_liter:.3f}".replace(".", ",") if r.price_per_liter else "",
                f"{r.liters:.3f}".replace(".", ",") if r.liters else "",
                f"{r.total_amount:.2f}".replace(".", ",") if r.total_amount else "",
                "Sim" if r.is_full_tank else "Não",
                f"{r.consumption_km_l:.2f}".replace(".", ",") if r.consumption_km_l else "",
                r.station_name,
                r.notes,
            ])

    @classmethod
    def _write_expenses_csv(cls, writer, user, vehicle, date_start, date_end):
        writer.writerow([
            "Data", "Veículo", "Categoria", "Tipo", "Odômetro (km)",
            "Valor (R$)", "Fornecedor / Local", "Observações"
        ])
        qs = Expense.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "category")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(occurred_at__gte=date_start)
        if date_end:
            qs = qs.filter(occurred_at__lte=date_end)

        for e in qs.order_by("-occurred_at"):
            writer.writerow([
                e.occurred_at.strftime("%d/%m/%Y"),
                e.vehicle.name,
                e.category.name if e.category else "",
                e.category.get_kind_display() if e.category else "",
                e.odometer if e.odometer is not None else "",
                f"{e.amount:.2f}".replace(".", ","),
                e.vendor_name,
                e.notes,
            ])

    @classmethod
    def _write_maintenance_csv(cls, writer, user, vehicle, date_start, date_end):
        writer.writerow([
            "Data", "Veículo", "Tipo de Serviço", "Natureza", "Odômetro (km)",
            "Valor Total (R$)", "Oficina / Mecânico", "Descrição", "Observações"
        ])
        qs = Maintenance.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "service_type")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(occurred_at__gte=date_start)
        if date_end:
            qs = qs.filter(occurred_at__lte=date_end)

        for m in qs.order_by("-occurred_at"):
            writer.writerow([
                m.occurred_at.strftime("%d/%m/%Y"),
                m.vehicle.name,
                m.service_type.name if m.service_type else "",
                m.get_maintenance_type_display(),
                m.odometer,
                f"{m.total_amount:.2f}".replace(".", ","),
                m.workshop_name,
                m.description,
                m.notes,
            ])

    @classmethod
    def _write_trips_csv(cls, writer, user, vehicle, date_start, date_end):
        writer.writerow([
            "Data/Hora Inicial", "Data/Hora Final", "Veículo", "Origem", "Destino",
            "Finalidade", "Odômetro Inicial", "Odômetro Final", "Distância (km)",
            "Valor/km (R$)", "Total (R$)", "Motorista", "Observações"
        ])
        qs = Trip.objects.filter(vehicle__owner_primary=user).select_related("vehicle")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(started_at__date__gte=date_start)
        if date_end:
            qs = qs.filter(started_at__date__lte=date_end)

        for t in qs.order_by("-started_at"):
            writer.writerow([
                t.started_at.strftime("%d/%m/%Y %H:%M"),
                t.ended_at.strftime("%d/%m/%Y %H:%M") if t.ended_at else "",
                t.vehicle.name,
                t.origin,
                t.destination,
                t.get_purpose_display(),
                t.start_odometer,
                t.end_odometer if t.end_odometer is not None else "",
                f"{t.distance:.1f}".replace(".", ",") if t.distance else "",
                f"{t.rate_per_km:.3f}".replace(".", ",") if t.rate_per_km else "",
                f"{t.total_cost:.2f}".replace(".", ",") if t.total_cost else "",
                t.driver_name,
                t.notes,
            ])

    @classmethod
    def _write_reminders_csv(cls, writer, user, vehicle):
        writer.writerow([
            "Veículo", "Título", "Tipo", "Data Limite", "Odômetro Limite (km)",
            "Status", "Recorrente", "Recorrência (Meses/Km)", "Observações"
        ])
        qs = Reminder.objects.filter(vehicle__owner_primary=user).select_related("vehicle")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)

        for r in qs.order_by("status", "due_date"):
            rec_str = ""
            if r.is_recurring:
                parts = []
                if r.recurrence_months:
                    parts.append(f"{r.recurrence_months}m")
                if r.recurrence_km:
                    parts.append(f"{r.recurrence_km}km")
                rec_str = " / ".join(parts)

            writer.writerow([
                r.vehicle.name,
                r.title,
                r.get_reminder_type_display(),
                r.due_date.strftime("%d/%m/%Y") if r.due_date else "",
                r.due_odometer if r.due_odometer else "",
                r.get_status_display(),
                "Sim" if r.is_recurring else "Não",
                rec_str,
                r.notes,
            ])

    @classmethod
    def _write_consolidated_csv(cls, writer, user, vehicle, date_start, date_end):
        writer.writerow([
            "Data", "Tipo de Registro", "Veículo", "Descrição / Item",
            "Categoria / Especificação", "Odômetro (km)", "Valor Total (R$)"
        ])
        records = []

        fuel_qs = Refueling.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "fuel_type")
        exp_qs = Expense.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "category")
        maint_qs = Maintenance.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "service_type")
        trips_qs = Trip.objects.filter(vehicle__owner_primary=user).select_related("vehicle")

        if vehicle:
            fuel_qs = fuel_qs.filter(vehicle=vehicle)
            exp_qs = exp_qs.filter(vehicle=vehicle)
            maint_qs = maint_qs.filter(vehicle=vehicle)
            trips_qs = trips_qs.filter(vehicle=vehicle)

        if date_start:
            fuel_qs = fuel_qs.filter(occurred_at__gte=date_start)
            exp_qs = exp_qs.filter(occurred_at__gte=date_start)
            maint_qs = maint_qs.filter(occurred_at__gte=date_start)
            trips_qs = trips_qs.filter(started_at__date__gte=date_start)

        if date_end:
            fuel_qs = fuel_qs.filter(occurred_at__lte=date_end)
            exp_qs = exp_qs.filter(occurred_at__lte=date_end)
            maint_qs = maint_qs.filter(occurred_at__lte=date_end)
            trips_qs = trips_qs.filter(started_at__date__lte=date_end)

        for f in fuel_qs:
            records.append({
                "date": f.occurred_at,
                "type": "Abastecimento",
                "vehicle": f.vehicle.name,
                "desc": f"{f.liters} L ({f.fuel_type.name if f.fuel_type else ''})",
                "spec": f.station_name,
                "odometer": f.odometer,
                "amount": f.total_amount,
            })
        for e in exp_qs:
            records.append({
                "date": e.occurred_at,
                "type": "Despesa",
                "vehicle": e.vehicle.name,
                "desc": e.category.name if e.category else "Outros",
                "spec": e.vendor_name,
                "odometer": e.odometer,
                "amount": e.amount,
            })
        for m in maint_qs:
            records.append({
                "date": m.occurred_at,
                "type": "Manutenção",
                "vehicle": m.vehicle.name,
                "desc": m.service_type.name if m.service_type else m.description,
                "spec": m.workshop_name,
                "odometer": m.odometer,
                "amount": m.total_amount,
            })
        for t in trips_qs:
            if t.total_cost:
                records.append({
                    "date": t.started_at.date(),
                    "type": "Viagem",
                    "vehicle": t.vehicle.name,
                    "desc": f"{t.origin} → {t.destination}",
                    "spec": t.get_purpose_display(),
                    "odometer": t.end_odometer or t.start_odometer,
                    "amount": t.total_cost,
                })

        records.sort(key=lambda x: x["date"], reverse=True)

        for r in records:
            writer.writerow([
                r["date"].strftime("%d/%m/%Y"),
                r["type"],
                r["vehicle"],
                r["desc"],
                r["spec"],
                r["odometer"] if r["odometer"] is not None else "",
                f"{r['amount']:.2f}".replace(".", ",") if r["amount"] else "0,00",
            ])

    # -------------------------------------------------------------------------
    # Excel (.xlsx) Export
    # -------------------------------------------------------------------------
    @classmethod
    def _export_excel(cls, user, entity, filename, vehicle, date_start, date_end) -> HttpResponse:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default blank sheet

        if entity == "consolidated":
            cls._build_excel_sheet_fuel(wb, user, vehicle, date_start, date_end)
            cls._build_excel_sheet_expenses(wb, user, vehicle, date_start, date_end)
            cls._build_excel_sheet_maintenance(wb, user, vehicle, date_start, date_end)
            cls._build_excel_sheet_trips(wb, user, vehicle, date_start, date_end)
            cls._build_excel_sheet_reminders(wb, user, vehicle)
        elif entity == "fuel":
            cls._build_excel_sheet_fuel(wb, user, vehicle, date_start, date_end)
        elif entity == "expenses":
            cls._build_excel_sheet_expenses(wb, user, vehicle, date_start, date_end)
        elif entity == "maintenance":
            cls._build_excel_sheet_maintenance(wb, user, vehicle, date_start, date_end)
        elif entity == "trips":
            cls._build_excel_sheet_trips(wb, user, vehicle, date_start, date_end)
        elif entity == "reminders":
            cls._build_excel_sheet_reminders(wb, user, vehicle)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        response = HttpResponse(
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @classmethod
    def _style_sheet(cls, ws, headers):
        ws.append(headers)
        for col_idx, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 25

    @classmethod
    def _autofit_columns(cls, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    @classmethod
    def _build_excel_sheet_fuel(cls, wb, user, vehicle, date_start, date_end):
        ws = wb.create_sheet(title="Abastecimentos")
        headers = [
            "Data", "Veículo", "Odômetro (km)", "Combustível", "Preço/L (R$)",
            "Volume (L)", "Valor Total (R$)", "Tanque Cheio", "Consumo (km/l)", "Posto", "Observações"
        ]
        cls._style_sheet(ws, headers)

        qs = Refueling.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "fuel_type")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(occurred_at__gte=date_start)
        if date_end:
            qs = qs.filter(occurred_at__lte=date_end)

        for r in qs.order_by("-occurred_at", "-odometer"):
            ws.append([
                r.occurred_at.strftime("%d/%m/%Y"),
                r.vehicle.name,
                r.odometer,
                r.fuel_type.name if r.fuel_type else "",
                float(r.price_per_liter) if r.price_per_liter else None,
                float(r.liters) if r.liters else None,
                float(r.total_amount) if r.total_amount else None,
                "Sim" if r.is_full_tank else "Não",
                float(r.consumption_km_l) if r.consumption_km_l else None,
                r.station_name,
                r.notes,
            ])
        cls._autofit_columns(ws)

    @classmethod
    def _build_excel_sheet_expenses(cls, wb, user, vehicle, date_start, date_end):
        ws = wb.create_sheet(title="Despesas")
        headers = [
            "Data", "Veículo", "Categoria", "Tipo", "Odômetro (km)",
            "Valor (R$)", "Fornecedor / Local", "Observações"
        ]
        cls._style_sheet(ws, headers)

        qs = Expense.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "category")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(occurred_at__gte=date_start)
        if date_end:
            qs = qs.filter(occurred_at__lte=date_end)

        for e in qs.order_by("-occurred_at"):
            ws.append([
                e.occurred_at.strftime("%d/%m/%Y"),
                e.vehicle.name,
                e.category.name if e.category else "",
                e.category.get_kind_display() if e.category else "",
                e.odometer,
                float(e.amount) if e.amount else None,
                e.vendor_name,
                e.notes,
            ])
        cls._autofit_columns(ws)

    @classmethod
    def _build_excel_sheet_maintenance(cls, wb, user, vehicle, date_start, date_end):
        ws = wb.create_sheet(title="Manutenções")
        headers = [
            "Data", "Veículo", "Tipo de Serviço", "Natureza", "Odômetro (km)",
            "Valor Total (R$)", "Oficina / Mecânico", "Descrição", "Observações"
        ]
        cls._style_sheet(ws, headers)

        qs = Maintenance.objects.filter(vehicle__owner_primary=user).select_related("vehicle", "service_type")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(occurred_at__gte=date_start)
        if date_end:
            qs = qs.filter(occurred_at__lte=date_end)

        for m in qs.order_by("-occurred_at"):
            ws.append([
                m.occurred_at.strftime("%d/%m/%Y"),
                m.vehicle.name,
                m.service_type.name if m.service_type else "",
                m.get_maintenance_type_display(),
                m.odometer,
                float(m.total_amount) if m.total_amount else None,
                m.workshop_name,
                m.description,
                m.notes,
            ])
        cls._autofit_columns(ws)

    @classmethod
    def _build_excel_sheet_trips(cls, wb, user, vehicle, date_start, date_end):
        ws = wb.create_sheet(title="Viagens")
        headers = [
            "Data/Hora Inicial", "Data/Hora Final", "Veículo", "Origem", "Destino",
            "Finalidade", "Odômetro Inicial", "Odômetro Final", "Distância (km)",
            "Valor/km (R$)", "Total (R$)", "Motorista", "Observações"
        ]
        cls._style_sheet(ws, headers)

        qs = Trip.objects.filter(vehicle__owner_primary=user).select_related("vehicle")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)
        if date_start:
            qs = qs.filter(started_at__date__gte=date_start)
        if date_end:
            qs = qs.filter(started_at__date__lte=date_end)

        for t in qs.order_by("-started_at"):
            ws.append([
                t.started_at.strftime("%d/%m/%Y %H:%M"),
                t.ended_at.strftime("%d/%m/%Y %H:%M") if t.ended_at else "",
                t.vehicle.name,
                t.origin,
                t.destination,
                t.get_purpose_display(),
                t.start_odometer,
                t.end_odometer,
                float(t.distance) if t.distance else None,
                float(t.rate_per_km) if t.rate_per_km else None,
                float(t.total_cost) if t.total_cost else None,
                t.driver_name,
                t.notes,
            ])
        cls._autofit_columns(ws)

    @classmethod
    def _build_excel_sheet_reminders(cls, wb, user, vehicle):
        ws = wb.create_sheet(title="Lembretes")
        headers = [
            "Veículo", "Título", "Tipo", "Data Limite", "Odômetro Limite (km)",
            "Status", "Recorrente", "Observações"
        ]
        cls._style_sheet(ws, headers)

        qs = Reminder.objects.filter(vehicle__owner_primary=user).select_related("vehicle")
        if vehicle:
            qs = qs.filter(vehicle=vehicle)

        for r in qs.order_by("status", "due_date"):
            ws.append([
                r.vehicle.name,
                r.title,
                r.get_reminder_type_display(),
                r.due_date.strftime("%d/%m/%Y") if r.due_date else "",
                r.due_odometer if r.due_odometer else "",
                r.get_status_display(),
                "Sim" if r.is_recurring else "Não",
                r.notes,
            ])
        cls._autofit_columns(ws)

